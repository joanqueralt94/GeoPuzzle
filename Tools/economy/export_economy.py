"""Export GeoQuizz economy values from economy_config.xlsx to economy_config.json.

Usage: python3 export_economy.py [economy_config.xlsx] [economy_config.json]
Reads the "Config" sheet (columns: key, value, type, ...) and writes a flat
key -> value JSON. Flat keys map one-to-one to Unity Remote Config / Firebase
Remote Config keys, so the same names work locally and remotely.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

CASTS = {"int": int, "float": float, "bool": lambda v: str(v).strip().lower() in ("true", "1", "yes"), "string": str}


def export(xlsx_path, json_path):
    ws = load_workbook(xlsx_path, data_only=True)["Config"]
    header = [c.value for c in ws[1]]
    col = {name: i for i, name in enumerate(header)}
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key = row[col["key"]]
        if not key or str(key).startswith("#"):
            continue
        value, kind = row[col["value"]], row[col["type"]]
        if key in out:
            raise ValueError(f"duplicate key: {key}")
        out[key] = CASTS[kind](value)
    Path(json_path).write_text(json.dumps(out, indent=2) + "\n")
    print(f"wrote {len(out)} keys to {json_path}")


if __name__ == "__main__":
    here = Path(__file__).parent
    export(sys.argv[1] if len(sys.argv) > 1 else here / "economy_config.xlsx",
           sys.argv[2] if len(sys.argv) > 2 else here / "economy_config.json")
